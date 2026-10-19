"""Deterministic evaluation of traceability and generated artifacts.

The evaluator deliberately has no application/database imports.  It accepts a
plain dictionary so it can run against a live-project snapshot, a CI fixture,
or an archived experiment result.
"""

from __future__ import annotations

import re
import shutil
import statistics
import subprocess
import tempfile
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

REQ_RE = re.compile(r"^REQ-\d{3,4}$", re.IGNORECASE)
STORY_RE = re.compile(r"^US-\d{3,4}$", re.IGNORECASE)
REFERENCE_RE = re.compile(r"\b(?:REQ|US)-\d{1,4}\b", re.IGNORECASE)
MERMAID_HEADER_RE = re.compile(r"^\s*(?:flowchart|graph)\s+(?:TD|TB|LR|RL)\b", re.IGNORECASE)

CANONICAL_SECTIONS = (
    "title",
    "stakeholders",
    "version_history",
    "reviews",
    "contents",
    "business_overview",
    "product_scope",
    "tech_ops",
    "appendix",
)
CONTENT_REQUIRED_SECTIONS = (
    "title",
    "stakeholders",
    "business_overview",
    "product_scope",
    "tech_ops",
    "appendix",
)
EXPERT_DIMENSIONS = (
    "requirements_complete",
    "traceability_correct",
    "acceptance_criteria_testable",
    "diagrams_match_requirements",
    "prd_and_exports_complete",
)


def _text(value: Any) -> str:
    return str(value or "").strip()


def _check(
    checks: list[dict[str, Any]],
    category: str,
    check_id: str,
    passed: bool | None,
    detail: str,
    *,
    artifact: str | None = None,
    critical: bool = False,
) -> None:
    checks.append(
        {
            "category": category,
            "check_id": check_id,
            "status": "skip" if passed is None else ("pass" if passed else "fail"),
            "artifact": artifact,
            "detail": detail,
            "critical": bool(critical and passed is False),
        }
    )


def _normalise(snapshot: dict[str, Any]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Return requirements and stories from nested or flat snapshot shapes."""
    requirements = list(snapshot.get("requirements") or [])
    top_stories = list(snapshot.get("user_stories") or [])
    stories: list[dict[str, Any]] = []
    for requirement in requirements:
        req_id = requirement.get("id")
        req_code = requirement.get("requirement_code")
        for story in requirement.get("user_stories") or []:
            item = dict(story)
            item.setdefault("requirement_id", req_id)
            item.setdefault("requirement_code", req_code)
            stories.append(item)
    stories.extend(top_stories)

    # The traceability endpoint returns rows instead of requirements.
    if not requirements and snapshot.get("rows") is not None:
        for row in snapshot.get("rows") or []:
            requirement = dict(row.get("requirement") or {})
            requirement["user_stories"] = row.get("user_stories") or []
            requirement["prd_sections"] = row.get("prd_sections") or []
            requirement["diagrams"] = row.get("diagrams") or []
            requirements.append(requirement)
        return _normalise({**snapshot, "requirements": requirements, "rows": None})
    return requirements, stories


def _summarise(checks: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = list(checks)
    passed = sum(row["status"] == "pass" for row in rows)
    failed = sum(row["status"] == "fail" for row in rows)
    skipped = sum(row["status"] == "skip" for row in rows)
    denominator = passed + failed
    return {
        "passed": passed,
        "failed": failed,
        "skipped": skipped,
        "applicable": denominator,
        "pass_rate": round((passed / denominator * 100), 2) if denominator else None,
    }


def _extract_file_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt", ".tex", ".mmd", ".mermaid"}:
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        return "\n".join(p.text for p in document.paragraphs)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if not reader.pages:
            raise ValueError("PDF contains no pages")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    raise ValueError(f"unsupported file type: {suffix or '<none>'}")


def _evaluate_files(files: list[Any], checks: list[dict[str, Any]], base_dir: Path) -> None:
    if not files:
        _check(checks, "generated_files", "exports_present", None, "No exported files supplied.")
        return
    for index, entry in enumerate(files, start=1):
        raw_path = entry.get("path") if isinstance(entry, dict) else entry
        path = Path(_text(raw_path))
        if not path.is_absolute():
            path = base_dir / path
        artifact = str(path)
        exists = path.is_file() and path.stat().st_size > 0
        _check(
            checks,
            "generated_files",
            f"file_{index}_exists",
            exists,
            "File exists and is non-empty." if exists else "File is missing or empty.",
            artifact=artifact,
            critical=True,
        )
        if not exists:
            continue
        try:
            content = _extract_file_text(path)
        except Exception as exc:
            _check(
                checks,
                "generated_files",
                f"file_{index}_opens",
                False,
                f"Could not open export: {exc}",
                artifact=artifact,
                critical=True,
            )
            continue
        _check(
            checks,
            "generated_files",
            f"file_{index}_opens",
            True,
            "Export opened successfully with an independent parser.",
            artifact=artifact,
        )
        _check(
            checks,
            "generated_files",
            f"file_{index}_has_text",
            bool(content.strip()),
            "Extracted readable text." if content.strip() else "Export has no readable text.",
            artifact=artifact,
            critical=True,
        )
        required_headings = entry.get("required_headings", []) if isinstance(entry, dict) else []
        for heading_index, heading in enumerate(required_headings, start=1):
            heading_text = _text(heading)
            present = heading_text.casefold() in content.casefold()
            _check(
                checks,
                "generated_files",
                f"file_{index}_heading_{heading_index}",
                present,
                f"Required heading found: {heading_text}" if present else f"Required heading missing: {heading_text}",
                artifact=artifact,
            )
        placeholders = sorted(
            token for token in ("TBD", "TODO", "{{", "[PLACEHOLDER]")
            if token.casefold() in content.casefold()
        )
        _check(
            checks,
            "generated_files",
            f"file_{index}_no_placeholders",
            not placeholders,
            f"Placeholder markers found: {placeholders}" if placeholders else "No common placeholder markers found.",
            artifact=artifact,
        )


def _evaluate_mermaid(diagrams: list[Any], checks: list[dict[str, Any]]) -> None:
    if not diagrams:
        _check(checks, "diagrams", "diagram_present", False, "No diagram supplied.", critical=True)
        return
    mmdc = shutil.which("mmdc")
    for index, entry in enumerate(diagrams, start=1):
        source = _text(entry.get("source") if isinstance(entry, dict) else entry)
        label = _text(entry.get("label")) if isinstance(entry, dict) else f"diagram-{index}"
        label = label or f"diagram-{index}"
        structurally_valid = bool(MERMAID_HEADER_RE.match(source))
        _check(
            checks,
            "diagrams",
            f"diagram_{index}_syntax",
            structurally_valid,
            "Supported Mermaid flowchart header found." if structurally_valid else "Missing or unsupported Mermaid flowchart header.",
            artifact=label,
            critical=True,
        )
        if not structurally_valid:
            continue
        if not mmdc:
            _check(
                checks,
                "diagrams",
                f"diagram_{index}_renders",
                None,
                "Mermaid CLI (mmdc) is not installed; render validation skipped.",
                artifact=label,
            )
            continue
        with tempfile.TemporaryDirectory(prefix="agenticdia-eval-") as tmp:
            source_path = Path(tmp) / "diagram.mmd"
            output_path = Path(tmp) / "diagram.svg"
            source_path.write_text(source, encoding="utf-8")
            result = subprocess.run(
                [mmdc, "-i", str(source_path), "-o", str(output_path)],
                capture_output=True,
                text=True,
                timeout=30,
                check=False,
            )
            rendered = result.returncode == 0 and output_path.is_file() and output_path.stat().st_size > 0
            detail = "Mermaid rendered successfully." if rendered else (result.stderr.strip() or "Mermaid render failed.")
            _check(
                checks,
                "diagrams",
                f"diagram_{index}_renders",
                rendered,
                detail[:500],
                artifact=label,
                critical=True,
            )


def _aggregate_expert_reviews(reviews: list[dict[str, Any]], threshold: float) -> dict[str, Any]:
    reviewers = {_text(row.get("reviewer_id")) for row in reviews if _text(row.get("reviewer_id"))}
    dimensions: dict[str, dict[str, Any]] = {}
    invalid: list[str] = []
    for dimension in EXPERT_DIMENSIONS:
        scores: list[float] = []
        for row in reviews:
            value = (row.get("scores") or {}).get(dimension)
            if value is None:
                continue
            try:
                score = float(value)
            except (TypeError, ValueError):
                invalid.append(f"{row.get('reviewer_id', 'unknown')}:{dimension}")
                continue
            if not 1 <= score <= 5:
                invalid.append(f"{row.get('reviewer_id', 'unknown')}:{dimension}")
                continue
            scores.append(score)
        dimensions[dimension] = {
            "scores": scores,
            "average": round(statistics.mean(scores), 2) if scores else None,
            "passed": bool(scores) and statistics.mean(scores) >= threshold,
        }
    complete = len(reviewers) >= 2 and all(len(item["scores"]) >= 2 for item in dimensions.values())
    averages = [item["average"] for item in dimensions.values() if item["average"] is not None]
    return {
        "reviewer_count": len(reviewers),
        "complete": complete and not invalid,
        "invalid_scores": invalid,
        "dimensions": dimensions,
        "overall_average": round(statistics.mean(averages), 2) if averages else None,
        "passed": complete and not invalid and all(item["passed"] for item in dimensions.values()),
    }


def _aggregate_time(runs: list[dict[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {"run_count": len(runs)}
    for key in ("generation_seconds", "acceptance_seconds", "revisions", "prompts"):
        values: list[float] = []
        for run in runs:
            try:
                value = float(run[key])
            except (KeyError, TypeError, ValueError):
                continue
            if value >= 0:
                values.append(value)
        result[key] = {
            "count": len(values),
            "average": round(statistics.mean(values), 2) if values else None,
            "median": round(statistics.median(values), 2) if values else None,
        }
    return result


def evaluate_snapshot(
    snapshot: dict[str, Any],
    *,
    base_dir: str | Path = ".",
    automated_threshold: float = 95.0,
    expert_threshold: float = 4.0,
) -> dict[str, Any]:
    """Evaluate one project snapshot and return a JSON-serialisable report."""
    checks: list[dict[str, Any]] = []
    requirements, stories = _normalise(snapshot)
    req_codes = [_text(row.get("requirement_code")).upper() for row in requirements]
    story_codes = [_text(row.get("ticket_code")).upper() for row in stories]
    known_codes = {code for code in req_codes + story_codes if code}

    for index, code in enumerate(req_codes, start=1):
        _check(checks, "artifact_ids", f"requirement_{index}_id", bool(REQ_RE.fullmatch(code)), f"Requirement ID: {code or '<blank>'}", artifact=code or None)
    for index, code in enumerate(story_codes, start=1):
        _check(checks, "artifact_ids", f"story_{index}_id", bool(STORY_RE.fullmatch(code)), f"Story ID: {code or '<blank>'}", artifact=code or None)
    duplicate_reqs = sorted(code for code, count in Counter(req_codes).items() if code and count > 1)
    duplicate_stories = sorted(code for code, count in Counter(story_codes).items() if code and count > 1)
    _check(checks, "artifact_ids", "requirement_ids_unique", not duplicate_reqs, f"Duplicate requirement IDs: {duplicate_reqs}" if duplicate_reqs else "Requirement IDs are unique.", critical=bool(duplicate_reqs))
    _check(checks, "artifact_ids", "story_ids_unique", not duplicate_stories, f"Duplicate story IDs: {duplicate_stories}" if duplicate_stories else "Story IDs are unique.", critical=bool(duplicate_stories))

    stories_by_req: dict[str, list[dict[str, Any]]] = {}
    for story in stories:
        parent = _text(story.get("requirement_code") or story.get("requirement_id"))
        stories_by_req.setdefault(parent, []).append(story)
    for index, requirement in enumerate(requirements, start=1):
        code = req_codes[index - 1]
        nested = requirement.get("user_stories") or []
        linked = nested or stories_by_req.get(code) or stories_by_req.get(_text(requirement.get("id"))) or []
        _check(checks, "traceability", f"{code or index}_has_story", bool(linked), "Requirement has a linked user story." if linked else "Requirement has no linked user story.", artifact=code, critical=not bool(linked))
        section_links = requirement.get("prd_sections")
        if section_links is not None:
            _check(checks, "traceability", f"{code or index}_has_prd_link", bool(section_links), "Requirement is referenced by a PRD section." if section_links else "Requirement has no PRD-section reference.", artifact=code)
    for index, story in enumerate(stories, start=1):
        code = story_codes[index - 1]
        criteria = [item for item in (story.get("acceptance_criteria") or []) if _text(item.get("criteria_text") if isinstance(item, dict) else item)]
        _check(checks, "traceability", f"{code or index}_has_criteria", bool(criteria), "Story has acceptance criteria." if criteria else "Story has no acceptance criteria.", artifact=code, critical=not bool(criteria))
        parent_code = _text(story.get("requirement_code")).upper()
        parent_id = _text(story.get("requirement_id"))
        known_parent_ids = {_text(row.get("id")) for row in requirements}
        parent_exists = bool(parent_code and parent_code in set(req_codes)) or bool(parent_id and parent_id in known_parent_ids)
        _check(checks, "traceability", f"{code or index}_parent_exists", parent_exists, "Story parent requirement exists." if parent_exists else "Story points to a missing requirement.", artifact=code, critical=not parent_exists)

    sections = list(snapshot.get("prd_sections") or snapshot.get("sections") or [])
    section_map = {_text(row.get("section_key")): row for row in sections if isinstance(row, dict)}
    for key in CANONICAL_SECTIONS:
        _check(checks, "prd", f"section_{key}_exists", key in section_map, "Canonical PRD section exists." if key in section_map else "Canonical PRD section is missing.", artifact=key)
    for key in CONTENT_REQUIRED_SECTIONS:
        if key in section_map:
            content = _text(section_map[key].get("content"))
            _check(checks, "prd", f"section_{key}_not_empty", bool(content), "Required generated section contains content." if content else "Required generated section is empty.", artifact=key)

    documents = list(snapshot.get("prd_documents") or [])
    _check(
        checks,
        "prd",
        "generated_prd_present",
        bool(documents),
        "At least one generated PRD document exists." if documents else "No generated PRD document supplied.",
        critical=not bool(documents),
    )
    if documents:
        latest = max(documents, key=lambda item: item.get("version", 0))
        body = _text(latest.get("prd_markdown"))
        _check(
            checks,
            "prd",
            "latest_prd_not_empty",
            bool(body),
            "Latest PRD document contains content." if body else "Latest PRD document is empty.",
            artifact=f"PRD v{latest.get('version', '?')}",
            critical=not bool(body),
        )

    # Traceability API diagram entries contain labels/versions but not source.
    # Prefer explicit source entries, otherwise use sources from PRD documents.
    diagrams = [
        item for item in (snapshot.get("diagrams") or [])
        if not isinstance(item, dict) or "source" in item
    ]
    if not diagrams:
        diagrams = [
            {"label": f"PRD v{doc.get('version', index)}", "source": doc.get("mermaid_diagram", "")}
            for index, doc in enumerate(documents, start=1)
            if _text(doc.get("mermaid_diagram"))
        ]
    reference_sources = [(f"section:{key}", _text(row.get("content"))) for key, row in section_map.items()]
    reference_sources.extend(
        (f"diagram:{index}", _text(item.get("source") if isinstance(item, dict) else item))
        for index, item in enumerate(diagrams, start=1)
    )
    stale = []
    for label, content in reference_sources:
        for reference in sorted({match.upper() for match in REFERENCE_RE.findall(content)}):
            if reference not in known_codes:
                stale.append({"reference": reference, "source": label})
    _check(checks, "references", "all_references_resolve", not stale, f"Stale references: {stale}" if stale else "All artifact references resolve.", critical=bool(stale))

    _evaluate_mermaid(diagrams, checks)
    _evaluate_files(list(snapshot.get("generated_files") or []), checks, Path(base_dir))

    categories = {
        category: _summarise(row for row in checks if row["category"] == category)
        for category in sorted({row["category"] for row in checks})
    }
    automated = _summarise(checks)
    critical_errors = [row for row in checks if row["critical"]]
    automated["threshold"] = automated_threshold
    automated["critical_error_count"] = len(critical_errors)
    automated["passed_threshold"] = bool(
        automated["pass_rate"] is not None
        and automated["pass_rate"] >= automated_threshold
        and not critical_errors
    )

    expert = _aggregate_expert_reviews(list(snapshot.get("expert_reviews") or []), expert_threshold)
    expert["threshold"] = expert_threshold
    timing = _aggregate_time(list(snapshot.get("runs") or []))
    return {
        "project_id": snapshot.get("project_id"),
        "thresholds": {
            "automated_pass_rate": automated_threshold,
            "expert_dimension_average": expert_threshold,
            "critical_errors_allowed": 0,
            "note": "These thresholds are defined for this project's evaluation.",
        },
        "automated": automated,
        "categories": categories,
        "critical_errors": critical_errors,
        "checks": checks,
        "expert_review": expert,
        "time_measurement": timing,
    }
